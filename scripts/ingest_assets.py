#!/usr/bin/env python3
"""Drop-folder ingest for ZIN assets.

1. Classify everything in assets/drop/
2. Move into assets/zin/{bundles,workbook,extracts,brief}
3. Rebuild PDF manifests
4. Project data/zin/v1/ for ./pakket

Usage (from repo root):
  python scripts/ingest_assets.py
  python scripts/ingest_assets.py --dry-run
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
DROP = ASSETS / "drop"
ZIN = ASSETS / "zin"
BUNDLES_INT = ZIN / "bundles" / "int"
BUNDLES_GVS = ZIN / "bundles" / "gvs"
WORKBOOK = ZIN / "workbook"
EXTRACTS = ZIN / "extracts"
BRIEF = ZIN / "brief"
META = ZIN / "meta"
OUT = ROOT / "data" / "zin" / "v1"

SCHEMA = 1
AUTHORITY = "zin"
AS_OF = "2026-09-24"

ZIN_EXTRA_ALLOW = {
    "traject",
    "therapeutic_area",
    "indication_group",
    "outcome",
    "icer_manufacturer",
    "icer_zin",
    "budget_impact_mln",
    "patients_per_year",
    "price_cut_pct",
    "orphan",
    "extraction_status",
    "source_url",
    "keywords",
    "quote",
}


def _log(msg: str) -> None:
    print(msg)


def _date(val: object) -> str | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date().isoformat()
    if isinstance(val, date):
        return val.isoformat()
    text = str(val).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return text[:10]


def _str(val: object) -> str:
    if val is None:
        return ""
    return str(val).strip()


def _bool_checked(val: object) -> bool:
    return _str(val).casefold() in {"ja", "yes", "true", "1"}


def _write_json(path: Path, obj: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(obj, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [json.dumps(r, ensure_ascii=False, sort_keys=True) for r in rows]
    path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")


def _zip_names(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as zf:
        return [n.replace("\\", "/") for n in zf.namelist()]


def classify_zip(path: Path) -> str:
    """Return kind: int_bundle | gvs_bundle | workbook | extracts | handover | unknown."""
    names = _zip_names(path)
    flat = " ".join(names).casefold()
    if any(re.match(r"^INT-\d+/", n) for n in names):
        return "int_bundle"
    if any(re.match(r"^GVS-\d+/", n) for n in names):
        return "gvs_bundle"
    if any(n.endswith("ZIN_pakketadviezen_database.xlsx") for n in names):
        return "handover"
    if path.suffix.casefold() == ".xlsx" or any(
        n.casefold().endswith(".xlsx") and "pakketadvies" in n.casefold() for n in names
    ):
        return "workbook"
    if any(re.match(r"^GVS-\d+__.+\.txt$", Path(n).name) for n in names):
        return "extracts"
    if "zin_text_extracts" in path.name.casefold() or "zin_text_extracts" in flat:
        return "extracts"
    if "project_brief.md" in flat or "gvs_kb_handover/" in flat or "files.zip" == path.name:
        return "handover"
    # Outer wrapper often only holds nested zip + md
    if any(n.casefold().endswith(".zip") for n in names) and any(
        "brief" in n.casefold() or "handover" in n.casefold() for n in names
    ):
        return "handover"
    return "unknown"


def _unique_dest(dest_dir: Path, name: str) -> Path:
    dest = dest_dir / name
    if not dest.exists():
        return dest
    stem, suf = Path(name).stem, Path(name).suffix
    i = 2
    while True:
        cand = dest_dir / f"{stem}_{i}{suf}"
        if not cand.exists():
            return cand
        i += 1


def place_file(src: Path, dest_dir: Path, *, dry_run: bool) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = _unique_dest(dest_dir, src.name)
    _log(f"  move {src.name} -> {dest.relative_to(ROOT)}")
    if not dry_run:
        shutil.move(str(src), str(dest))
    return dest


def unpack_handover(path: Path, *, dry_run: bool) -> None:
    """Pull workbook / extracts / brief out of a handover or files.zip."""
    with tempfile.TemporaryDirectory(prefix="pakket-drop-") as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(path) as zf:
            zf.extractall(tmp_path)
        # Nested handover zip
        for nested in tmp_path.rglob("*.zip"):
            if nested.name.startswith("ZIN_KB_handover") or "handover" in nested.name.casefold():
                with zipfile.ZipFile(nested) as zf:
                    zf.extractall(tmp_path / "_handover")
        # Workbook
        for xlsx in tmp_path.rglob("ZIN_pakketadviezen_database.xlsx"):
            dest = WORKBOOK / xlsx.name
            _log(f"  extract workbook -> {dest.relative_to(ROOT)}")
            if not dry_run:
                WORKBOOK.mkdir(parents=True, exist_ok=True)
                shutil.copy2(xlsx, dest)
        for name in ("codeboek.json", "schema_columns.json", "index_dossiers.csv"):
            hits = list(tmp_path.rglob(name))
            if hits and not dry_run:
                WORKBOOK.mkdir(parents=True, exist_ok=True)
                shutil.copy2(hits[0], WORKBOOK / name)
                _log(f"  extract {name}")
        # Extracts
        for ext in tmp_path.rglob("zin_text_extracts*.zip"):
            dest = EXTRACTS / ext.name
            _log(f"  extract texts -> {dest.relative_to(ROOT)}")
            if not dry_run:
                EXTRACTS.mkdir(parents=True, exist_ok=True)
                shutil.copy2(ext, dest)
        # Brief
        for name in ("PROJECT_BRIEF.md", "STATE.md", "STATE_remaining.txt"):
            hits = list(tmp_path.rglob(name))
            if hits:
                _log(f"  extract {name}")
                if not dry_run:
                    BRIEF.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(hits[0], BRIEF / name)
    if not dry_run:
        path.unlink(missing_ok=True)
        _log(f"  removed wrapper {path.name}")


def process_drop(*, dry_run: bool) -> list[str]:
    DROP.mkdir(parents=True, exist_ok=True)
    actions: list[str] = []
    items = sorted(
        p for p in DROP.iterdir() if p.is_file() and p.name not in {".gitkeep", ".DS_Store"}
    )
    if not items:
        _log("drop/: empty")
        return actions
    for path in items:
        if path.suffix.casefold() == ".xlsx":
            kind = "workbook"
        elif path.suffix.casefold() != ".zip":
            _log(f"skip (not zip/xlsx): {path.name}")
            continue
        else:
            kind = classify_zip(path)
        _log(f"classify {path.name} -> {kind}")
        actions.append(f"{path.name}:{kind}")
        if kind == "int_bundle":
            place_file(path, BUNDLES_INT, dry_run=dry_run)
        elif kind == "gvs_bundle":
            place_file(path, BUNDLES_GVS, dry_run=dry_run)
        elif kind == "workbook":
            place_file(path, WORKBOOK, dry_run=dry_run)
        elif kind == "extracts":
            place_file(path, EXTRACTS, dry_run=dry_run)
        elif kind == "handover":
            unpack_handover(path, dry_run=dry_run)
        else:
            _log(f"  leave in drop/ (unknown): {path.name}")
    return actions


def _scan_zip(zpath: Path, prefix: str) -> dict[str, dict]:
    dossiers: dict[str, dict] = {}
    with zipfile.ZipFile(zpath) as zf:
        for name in zf.namelist():
            norm = name.replace("\\", "/")
            if norm.endswith("/"):
                continue
            m = re.match(rf"^({prefix}-\d+)/(.+)$", norm)
            if not m:
                continue
            dossier, rest = m.group(1), m.group(2)
            rec = dossiers.setdefault(
                dossier,
                {
                    "id": dossier,
                    "zip": zpath.name,
                    "zips": [zpath.name],
                    "pdfs": [],
                    "bron_url": None,
                    "members": [],
                    "duplicate": False,
                },
            )
            if zpath.name not in rec["zips"]:
                rec["zips"].append(zpath.name)
                rec["duplicate"] = True
            rec["members"].append(rest)
            if Path(rest).name.lower() == "bron.txt":
                raw = zf.read(name).decode("utf-8-sig", errors="replace").strip()
                rec["bron_url"] = raw.splitlines()[0].strip() if raw else None
            elif rest.lower().endswith(".pdf"):
                rec["pdfs"].append(rest)
    for rec in dossiers.values():
        rec["pdfs"] = sorted(set(rec["pdfs"]))
        rec["members"] = sorted(set(rec["members"]))
        rec["zips"] = sorted(set(rec["zips"]))
        rec["zip"] = rec["zips"][0]
    return dossiers


def build_manifests() -> None:
    int_d: dict[str, dict] = {}
    for zpath in sorted(BUNDLES_INT.glob("*.zip")):
        part = _scan_zip(zpath, "INT")
        for did, rec in part.items():
            if did in int_d:
                prev = int_d[did]
                int_d[did] = {
                    **prev,
                    "pdfs": sorted(set(prev["pdfs"]) | set(rec["pdfs"])),
                    "members": sorted(set(prev["members"]) | set(rec["members"])),
                    "zips": sorted(set(prev["zips"]) | set(rec["zips"])),
                    "duplicate": True,
                    "zip": sorted(set(prev["zips"]) | set(rec["zips"]))[0],
                }
            else:
                int_d[did] = rec
    int_ids = sorted(int_d, key=lambda x: int(x.split("-")[1]))
    _write_json(
        META / "int_bundle_manifest.json",
        {
            "as_of": AS_OF,
            "track": "Intramuraal",
            "track_nl": "Sluis / ziekenhuisvergoeding",
            "bundle_dir": "assets/zin/bundles/int",
            "dossier_count": len(int_ids),
            "dossiers": {i: int_d[i] for i in int_ids},
        },
    )
    _log(f"manifest INT: {len(int_ids)} dossiers")

    gvs: dict[str, dict] = {}
    for zpath in sorted(BUNDLES_GVS.glob("*.zip")):
        for did, rec in _scan_zip(zpath, "GVS").items():
            gvs[did] = rec
    gvs_ids = sorted(gvs, key=lambda x: int(x.split("-")[1]))
    _write_json(
        META / "gvs_bundle_manifest.json",
        {
            "as_of": AS_OF,
            "track": "GVS",
            "track_nl": "Geneesmiddelenvergoedingssysteem (extramuraal)",
            "bundle_dir": "assets/zin/bundles/gvs",
            "dossier_count": len(gvs_ids),
            "dossiers": {i: gvs[i] for i in gvs_ids},
        },
    )
    _log(f"manifest GVS: {len(gvs_ids)} dossiers")


def _split_drug(geneesmiddel: str) -> tuple[str, str]:
    text = geneesmiddel.strip()
    m = re.match(r"^(.+?)\s*\(([^)]+)\)\s*$", text)
    if not m:
        return text, ""
    return m.group(1).strip(), m.group(2).replace("®", "").strip()


def _resolve_xlsx() -> Path:
    preferred = WORKBOOK / "ZIN_pakketadviezen_database.xlsx"
    if preferred.is_file():
        return preferred
    hits = sorted(WORKBOOK.glob("*.xlsx"))
    if hits:
        return hits[0]
    raise FileNotFoundError(f"no workbook in {WORKBOOK}")


def project_corpus() -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("openpyxl required: pip install openpyxl") from exc

    xlsx = _resolve_xlsx()
    _log(f"project from {xlsx.relative_to(ROOT)}")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    # Arguments
    ws = wb["Argumenten"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) if c is not None else "" for c in rows[0]]
    idx = {name: i for i, name in enumerate(header)}
    args: list[dict] = []
    for row in rows[1:]:
        native = _str(row[idx["Argument-ID"]])
        if not native:
            continue
        drug = _str(row[idx["Geneesmiddel"]])
        substance, brand = _split_drug(drug)
        extra = {
            "traject": _str(row[idx["Traject"]]),
            "therapeutic_area": _str(row[idx["Therapeutisch gebied"]]),
            "indication_group": _str(row[idx["Indicatiegroep (ziekte)"]]),
            "keywords": _str(row[idx["Trefwoorden"]]),
            "quote": _str(row[idx["Citaat (kort, optioneel)"]]),
        }
        extra = {k: v for k, v in extra.items() if v}
        bad = set(extra) - ZIN_EXTRA_ALLOW
        if bad:
            raise SystemExit(f"forbidden extra keys: {bad}")
        args.append(
            {
                "schema_version": SCHEMA,
                "authority": AUTHORITY,
                "id": f"{AUTHORITY}:{native}",
                "native_id": native,
                "dossier_id": _str(row[idx["Dossier-ID"]]),
                "advice_date": _date(row[idx["Datum advies"]]),
                "substance": substance,
                "brand": brand,
                "indication": _str(row[idx["Indicatiegroep (ziekte)"]]),
                "line": _str(row[idx["Behandellijn"]]),
                "criterion": _str(row[idx["Pakketcriterium"]]),
                "theme": _str(row[idx["Thema"]]),
                "weight": _str(row[idx["Gewicht in advies"]]),
                "position": _str(row[idx["Positie ZIN t.o.v. claim fabrikant"]]),
                "text": _str(row[idx["Argument ZIN (parafrase in eigen woorden)"]]),
                "precedent": _str(row[idx["Precedentwaarde / herbruikbaar voor"]]),
                "source_label": _str(row[idx["Brondocument"]]),
                "source_locator": _str(row[idx["Pagina / sectie"]]),
                "checked": _bool_checked(row[idx["Gecontroleerd"]]),
                "extra": extra,
            }
        )

    dossiers: list[dict] = []
    for sheet, traject in (("Intramuraal", "Intramuraal"), ("GVS", "GVS")):
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(c) if c is not None else "" for c in rows[0]]
        idx = {name: i for i, name in enumerate(header)}
        for row in rows[1:]:
            native = _str(row[idx["ID"]])
            if not native:
                continue
            extra = {
                "traject": traject,
                "therapeutic_area": _str(row[idx["Therapeutisch gebied*"]]),
                "indication_group": _str(row[idx["Indicatiegroep (ziekte)*"]]),
            }
            if "Uitkomst advies" in idx:
                extra["outcome"] = _str(row[idx["Uitkomst advies"]])
            if "Status extractie" in idx:
                extra["extraction_status"] = _str(row[idx["Status extractie"]])
            if "URL documentpagina ZIN" in idx:
                extra["source_url"] = _str(row[idx["URL documentpagina ZIN"]])
            if "Weesgeneesmiddel" in idx:
                extra["orphan"] = _str(row[idx["Weesgeneesmiddel"]])
            for key, col in (
                ("icer_manufacturer", "ICER fabrikant (€/QALY)"),
                ("icer_zin", "ICER ZIN basecase (€/QALY)"),
                ("budget_impact_mln", "Budgetimpact ZIN (€ mln/jaar)"),
                ("patients_per_year", "Aantal patiënten/jaar (ZIN)"),
                ("price_cut_pct", "Geadviseerde prijsreductie (%)"),
            ):
                if col in idx and row[idx[col]] not in (None, ""):
                    extra[key] = row[idx[col]]
            extra = {k: v for k, v in extra.items() if v not in ("", None)}
            bad = set(extra) - ZIN_EXTRA_ALLOW
            if bad:
                raise SystemExit(f"forbidden extra on dossier: {bad}")
            dossiers.append(
                {
                    "schema_version": SCHEMA,
                    "authority": AUTHORITY,
                    "id": f"{AUTHORITY}:{native}",
                    "native_id": native,
                    "advice_date": _date(row[idx["Datum advies"]]),
                    "substance": _str(row[idx["Stofnaam (hoofd)*"]]),
                    "brand": _str(row[idx["Merknaam*"]]),
                    "indication": _str(row[idx["Indicatie (uit titel)*"]]),
                    "line": _str(row[idx["Behandellijn*"]]),
                    "status": _str(row[idx["Status"]]),
                    "title": _str(row[idx["Titel advies (ZIN)"]]),
                    "extra": extra,
                }
            )
    wb.close()

    sources: list[dict] = []
    seen: set[str] = set()

    def add_src(row: dict) -> None:
        if row["id"] in seen:
            return
        seen.add(row["id"])
        sources.append(row)

    int_man = META / "int_bundle_manifest.json"
    gvs_man = META / "gvs_bundle_manifest.json"
    if int_man.is_file():
        for dossier, rec in json.loads(int_man.read_text())["dossiers"].items():
            zip_name = (rec.get("zips") or [rec.get("zip", "")])[0]
            for pdf in rec.get("pdfs", []):
                add_src(
                    {
                        "schema_version": SCHEMA,
                        "authority": AUTHORITY,
                        "id": f"{AUTHORITY}:pdf:{dossier}:{pdf}",
                        "dossier_id": dossier,
                        "path": f"assets/zin/bundles/int/{zip_name}#{dossier}/{pdf}",
                        "text": "",
                    }
                )
    if gvs_man.is_file():
        for dossier, rec in json.loads(gvs_man.read_text())["dossiers"].items():
            zip_name = rec.get("zip", "")
            for pdf in rec.get("pdfs", []):
                add_src(
                    {
                        "schema_version": SCHEMA,
                        "authority": AUTHORITY,
                        "id": f"{AUTHORITY}:pdf:{dossier}:{pdf}",
                        "dossier_id": dossier,
                        "path": f"assets/zin/bundles/gvs/{zip_name}#{dossier}/{pdf}",
                        "text": "",
                    }
                )
    for ext_zip in sorted(EXTRACTS.glob("*.zip")):
        with zipfile.ZipFile(ext_zip) as zf:
            for name in zf.namelist():
                if not name.lower().endswith(".txt"):
                    continue
                m = re.match(r"^(GVS-\d+)__", Path(name).name)
                if not m:
                    continue
                dossier = m.group(1)
                add_src(
                    {
                        "schema_version": SCHEMA,
                        "authority": AUTHORITY,
                        "id": f"{AUTHORITY}:txt:{dossier}:{Path(name).name}",
                        "dossier_id": dossier,
                        "path": f"assets/zin/extracts/{ext_zip.name}#{name}",
                        "text": "",
                    }
                )

    OUT.mkdir(parents=True, exist_ok=True)
    _write_json(
        OUT / "authority.json",
        {
            "id": AUTHORITY,
            "name": "Zorginstituut Nederland",
            "as_of": AS_OF,
            "schema_version": SCHEMA,
            "reserved": False,
        },
    )
    _write_json(
        OUT / "allowlist.json",
        {"authority": AUTHORITY, "extra_keys": sorted(ZIN_EXTRA_ALLOW)},
    )
    _write_json(
        OUT / "codebook.json",
        {
            "authority": AUTHORITY,
            "weight_order": ["Doorslaggevend", "Ondersteunend", "Zijdelings"],
        },
    )
    _write_jsonl(OUT / "arguments.jsonl", args)
    _write_jsonl(OUT / "dossiers.jsonl", dossiers)
    _write_jsonl(OUT / "sources.jsonl", sources)
    _log(
        f"projected {len(args)} arguments, {len(dossiers)} dossiers, "
        f"{len(sources)} sources -> {OUT.relative_to(ROOT)}"
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--skip-drop", action="store_true", help="only rebuild + project")
    args = parser.parse_args()
    for d in (DROP, BUNDLES_INT, BUNDLES_GVS, WORKBOOK, EXTRACTS, BRIEF, META):
        d.mkdir(parents=True, exist_ok=True)
    if not args.skip_drop:
        process_drop(dry_run=args.dry_run)
    if args.dry_run:
        _log("dry-run: skip manifests and project")
        return 0
    build_manifests()
    project_corpus()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
